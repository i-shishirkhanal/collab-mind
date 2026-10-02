"""Builds small real documents on the fly so tests don't depend on committed
binary fixtures."""

import io

import pytest


@pytest.fixture
def pdf_bytes():
    from reportlab.pdfgen import canvas

    def make(pages: list[str]) -> bytes:
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        for text in pages:
            if text:
                c.drawString(72, 720, text)
            c.showPage()
        c.save()
        return buf.getvalue()

    return make


@pytest.fixture
def docx_bytes():
    import docx

    d = docx.Document()
    d.add_heading("Research Notes", 1)
    d.add_paragraph("Methods: we measured chlorophyll fluorescence.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Group", "Yield"
    table.cell(1, 0).text, table.cell(1, 1).text = "Control", "4.2"
    d.add_heading("Results", 2)
    d.add_paragraph("Yield increased 12 percent.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def pptx_bytes():
    from pptx import Presentation

    p = Presentation()
    for title, body in [("Intro", "Why photosynthesis matters"), ("Findings", "Light intensity drives yield")]:
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


@pytest.fixture
def xlsx_bytes():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yield"
    ws.append(["Group", "Yield"])
    ws.append(["Control", 4.2])
    ws.append(["Treated", 4.7])
    wb.create_sheet("Notes").append(["Sample size", "30"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
